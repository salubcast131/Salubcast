from .shared import *

@app.route("/companies", methods=["GET", "POST"])
@superadmin_required
def companies_manager() -> str:
    pending_delete_company_id = request.args.get("confirm_delete", "").strip()
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "create":
            name = request.form.get("name", "").strip()
            slug = request.form.get("slug", "").strip().lower().replace(" ", "-")
            if not name:
                flash("Bedrijfsnaam ontbreekt.")
                return redirect(url_for("companies_manager"))
            if fetch_one("SELECT id FROM companies WHERE lower(name) = lower(?) LIMIT 1", (name,)):
                flash("Er bestaat al een bedrijf met deze naam.")
                return redirect(url_for("companies_manager"))
            if not slug:
                slug = name.lower().replace(" ", "-")
            if fetch_one("SELECT id FROM companies WHERE slug = ? LIMIT 1", (slug,)):
                flash("Deze slug bestaat al.")
                return redirect(url_for("companies_manager"))
            trial_days = parse_int(os.environ.get('SALUBCAST_DEFAULT_TRIAL_DAYS', '14'), 14, minimum=1, maximum=365)
            trial_ends_at = (datetime.now(timezone.utc) + timedelta(days=trial_days)).isoformat()
            execute("INSERT INTO companies (id, name, slug, is_active, plan_name, billing_status, trial_ends_at, created_at) VALUES (?, ?, ?, 1, ?, ?, ?, ?)", (str(uuid.uuid4()), name, slug, "starter", "trial", trial_ends_at, now_iso()))
            flash("Bedrijf aangemaakt.")
            return redirect(url_for("companies_manager"))
        if action == "toggle_active":
            company_id = request.form.get("company_id", "")
            company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (company_id,))
            if company:
                execute("UPDATE companies SET is_active = ? WHERE id = ?", (0 if company['is_active'] == 1 else 1, company_id))
                flash("Bedrijfsstatus aangepast.")
            return redirect(url_for("companies_manager"))
        if action == "update_subscription":
            company_id = request.form.get("company_id", "")
            company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (company_id,))
            if not company:
                flash("Bedrijf niet gevonden.")
                return redirect(url_for("companies_manager"))
            plan_name = request.form.get("plan_name", "starter").strip().lower()
            billing_status = request.form.get("billing_status", "trial").strip().lower()
            if plan_name not in {"starter", "professional", "enterprise"}:
                plan_name = "starter"
            if billing_status not in {"trial", "active", "past_due", "unpaid", "canceled"}:
                billing_status = "trial"
            execute(
                "UPDATE companies SET plan_name = ?, billing_status = ? WHERE id = ?",
                (plan_name, billing_status, company_id),
            )
            flash("Abonnement bijgewerkt.")
            return redirect(url_for("companies_manager"))
        if action == "delete_company":
            company_id = request.form.get("company_id", "")
            return redirect(url_for("companies_manager", confirm_delete=company_id))
        if action == "confirm_delete_company":
            company_id = request.form.get("company_id", "")
            active_company_id = current_company_id()
            if company_id == active_company_id:
                flash("Je kunt het huidige sessiebedrijf niet verwijderen.")
                return redirect(url_for("companies_manager"))
            company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (company_id,))
            if not company:
                flash("Bedrijf niet gevonden.")
                return redirect(url_for("companies_manager"))
            media_rows = fetch_all("SELECT filename FROM media WHERE company_id = ?", (company_id,))
            media_ids = [row["id"] for row in fetch_all("SELECT id FROM media WHERE company_id = ?", (company_id,))]
            logo_filename = company["logo_filename"] if "logo_filename" in company.keys() else None
            try:
                conn = db()
                try:
                    if media_ids:
                        placeholders = ",".join("?" for _ in media_ids)
                        conn.execute(f"DELETE FROM playlist_items WHERE media_id IN ({placeholders})", tuple(media_ids))
                    conn.execute("DELETE FROM feed_items WHERE feed_id IN (SELECT id FROM feeds WHERE company_id = ?)", (company_id,))
                    conn.execute("DELETE FROM feeds WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM playlist_items WHERE playlist_id IN (SELECT id FROM playlists WHERE company_id = ?)", (company_id,))
                    conn.execute("DELETE FROM schedules WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM screens WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM playlists WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM media WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM users WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM audit_logs WHERE company_id = ?", (company_id,))
                    conn.execute("DELETE FROM companies WHERE id = ?", (company_id,))
                    conn.commit()
                finally:
                    conn.close()
            except sqlite3.OperationalError:
                flash("Database is nog bezig. Probeer het bedrijf over een paar seconden opnieuw te verwijderen.")
                return redirect(url_for("companies_manager"))
            for media_row in media_rows:
                filename = str(media_row["filename"] or "").strip()
                if not filename:
                    continue
                path = UPLOAD_DIR / filename
                if path.exists():
                    try:
                        path.unlink()
                    except Exception:
                        pass
            if logo_filename:
                logo_path = company_logo_path(logo_filename)
                if logo_path.exists():
                    try:
                        logo_path.unlink()
                    except Exception:
                        pass
            flash("Bedrijf verwijderd.")
            return redirect(url_for("companies_manager"))

    companies = fetch_all("SELECT * FROM companies ORDER BY created_at DESC")
    delete_preview = None
    if pending_delete_company_id:
        preview_company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (pending_delete_company_id,))
        if preview_company and preview_company["id"] != current_company_id():
            delete_preview = {
                "company": preview_company,
                "users": fetch_one("SELECT COUNT(*) AS c FROM users WHERE company_id = ?", (pending_delete_company_id,))["c"],
                "screens": fetch_one("SELECT COUNT(*) AS c FROM screens WHERE company_id = ?", (pending_delete_company_id,))["c"],
                "playlists": fetch_one("SELECT COUNT(*) AS c FROM playlists WHERE company_id = ?", (pending_delete_company_id,))["c"],
                "media": fetch_one("SELECT COUNT(*) AS c FROM media WHERE company_id = ?", (pending_delete_company_id,))["c"],
                "feeds": fetch_one("SELECT COUNT(*) AS c FROM feeds WHERE company_id = ?", (pending_delete_company_id,))["c"],
                "schedules": fetch_one("SELECT COUNT(*) AS c FROM schedules WHERE company_id = ?", (pending_delete_company_id,))["c"],
            }
    content = render_template_string(
        """
        {% if delete_preview %}
        <div class="card" style="border-color:#ef4444; box-shadow:0 14px 36px rgba(127,29,29,.22); margin-bottom:16px;">
          <h2>Bevestig verwijderen</h2>
          <p class="muted">Je staat op het punt bedrijf <strong>{{ delete_preview.company['name'] }}</strong> te verwijderen. Dit wist ook alle gekoppelde data hieronder.</p>
          <div class="grid cols-4" style="margin:14px 0;">
            <div class="card"><div class="muted">Gebruikers</div><div class="stat">{{ delete_preview.users }}</div></div>
            <div class="card"><div class="muted">Schermen</div><div class="stat">{{ delete_preview.screens }}</div></div>
            <div class="card"><div class="muted">Playlists</div><div class="stat">{{ delete_preview.playlists }}</div></div>
            <div class="card"><div class="muted">Media</div><div class="stat">{{ delete_preview.media }}</div></div>
            <div class="card"><div class="muted">Feeds</div><div class="stat">{{ delete_preview.feeds }}</div></div>
            <div class="card"><div class="muted">Schema's</div><div class="stat">{{ delete_preview.schedules }}</div></div>
          </div>
          <div class="inline">
            <form method="post">
              <input type="hidden" name="action" value="confirm_delete_company">
              <input type="hidden" name="company_id" value="{{ delete_preview.company['id'] }}">
              <button class="danger" type="submit" style="width:auto;">Definitief verwijderen</button>
            </form>
            <a href="{{ url_for('companies_manager') }}"><button class="secondary" type="button" style="width:auto;">Annuleren</button></a>
          </div>
        </div>
        {% endif %}
        <div class="grid two">
          <div class="card">
            <h2>Nieuw bedrijf</h2>
            <form method="post">
              <input type="hidden" name="action" value="create">
              <input name="name" placeholder="Bedrijfsnaam">
              <input name="slug" placeholder="Slug, bijv. newconsultancy">
              <button type="submit">Bedrijf maken</button>
            </form>
            <p class="muted">Nieuwe bedrijven starten standaard in trial op het starter-plan.</p>
          </div>
          <div class="card">
            <h2>Bedrijven</h2>
            <table class="table">
              <thead><tr><th>Naam</th><th>Slug</th><th>Status</th><th>Abonnement</th><th>Acties</th></tr></thead>
              <tbody>
                {% for company in companies %}
                <tr>
                  <td>{{ company['name'] }}<br><span class="muted">{{ company['billing_email'] or '-' }}</span></td>
                  <td>{{ company['slug'] }}</td>
                  <td><span class="pill {{ '' if company['is_active'] == 1 else 'off' }}">{{ 'Actief' if company['is_active'] == 1 else 'Geblokkeerd' }}</span></td>
                  <td>
                    <form method="post" style="display:grid; gap:8px;">
                      <input type="hidden" name="action" value="update_subscription">
                      <input type="hidden" name="company_id" value="{{ company['id'] }}">
                      <select name="plan_name">
                        <option value="starter" {{ 'selected' if (company['plan_name'] or 'starter') == 'starter' else '' }}>Starter</option>
                        <option value="professional" {{ 'selected' if (company['plan_name'] or 'starter') == 'professional' else '' }}>Professional</option>
                        <option value="enterprise" {{ 'selected' if (company['plan_name'] or 'starter') == 'enterprise' else '' }}>Enterprise</option>
                      </select>
                      <select name="billing_status">
                        <option value="trial" {{ 'selected' if (company['billing_status'] or 'trial') == 'trial' else '' }}>Trial</option>
                        <option value="active" {{ 'selected' if (company['billing_status'] or 'trial') == 'active' else '' }}>Active</option>
                        <option value="past_due" {{ 'selected' if (company['billing_status'] or 'trial') == 'past_due' else '' }}>Past due</option>
                        <option value="unpaid" {{ 'selected' if (company['billing_status'] or 'trial') == 'unpaid' else '' }}>Unpaid</option>
                        <option value="canceled" {{ 'selected' if (company['billing_status'] or 'trial') == 'canceled' else '' }}>Canceled</option>
                      </select>
                      <button type="submit" class="secondary" style="width:auto;">Abonnement opslaan</button>
                    </form>
                    <div class="muted" style="margin-top:8px;">{% if company['trial_ends_at'] %}Trial tot {{ company['trial_ends_at'][:10] }}{% endif %}</div>
                  </td>
                  <td><div class="inline"><form method="post"><input type="hidden" name="action" value="toggle_active"><input type="hidden" name="company_id" value="{{ company['id'] }}"><button class="secondary" style="width:auto;">{{ 'Blokkeer' if company['is_active'] == 1 else 'Activeer' }}</button></form>{% if company['id'] != session.get('company_id') %}<form method="post"><input type="hidden" name="action" value="delete_company"><input type="hidden" name="company_id" value="{{ company['id'] }}"><button class="danger" style="width:auto;">Delete</button></form>{% endif %}</div></td>
                </tr>
                {% endfor %}
              </tbody>
            </table>
          </div>
        </div>
        """,
        companies=companies,
        delete_preview=delete_preview,
    )
    return render_shell("Bedrijven", content)

@app.route("/users", methods=["GET", "POST"])
@admin_required
def users_manager() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "create_user")
        if action == "create_user":
            full_name = request.form.get("full_name", "").strip()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "").strip()
            role = request.form.get("role", "user").strip().lower()
            target_company_id = company_id
            if is_superadmin():
                target_company_id = request.form.get("company_id", company_id)
                allowed_roles = {"user", "company_admin", "superadmin"}
                if role not in allowed_roles:
                    role = "user"
            else:
                role = "user"
            if not full_name or not email or not password:
                flash("Naam, e-mail en wachtwoord zijn verplicht.")
                return redirect(url_for("users_manager"))
            if fetch_one("SELECT id FROM users WHERE email = ? LIMIT 1", (email,)):
                flash("Er bestaat al een account met dit e-mailadres.")
                return redirect(url_for("users_manager"))
            salt = secrets.token_hex(8)
            execute(
                "INSERT INTO users (id, company_id, email, password_hash, salt, full_name, role, is_active, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)",
                (str(uuid.uuid4()), target_company_id, email, hash_password(password, salt), salt, full_name, role, now_iso()),
            )
            log_event(actor_label(), "user_created", "user", email, full_name, target_company_id)
            flash("Gebruiker aangemaakt.")
            return redirect(url_for("users_manager"))
        if action == "toggle_active":
            user_id = request.form.get("user_id", "")
            user = fetch_one("SELECT * FROM users WHERE id = ? LIMIT 1", (user_id,)) if is_superadmin() else fetch_one("SELECT * FROM users WHERE id = ? AND company_id = ? LIMIT 1", (user_id, company_id))
            if user:
                execute("UPDATE users SET is_active = ? WHERE id = ?", (0 if user['is_active'] == 1 else 1, user_id))
                flash("Accountstatus aangepast.")
            return redirect(url_for("users_manager"))
        if action == "delete_user":
            user_id = request.form.get("user_id", "")
            if user_id == session.get("user_id"):
                flash("Je kunt jezelf niet verwijderen.")
                return redirect(url_for("users_manager"))
            if is_superadmin():
                execute("DELETE FROM users WHERE id = ?", (user_id,))
            else:
                execute("DELETE FROM users WHERE id = ? AND company_id = ?", (user_id, company_id))
            flash("Gebruiker verwijderd.")
            return redirect(url_for("users_manager"))
        if action == "change_role":
            if not is_superadmin():
                flash("Alleen superadmins mogen rollen wijzigen.")
                return redirect(url_for("users_manager"))
            user_id = request.form.get("user_id", "")
            role = request.form.get("role", "user").strip().lower()
            allowed_roles = {"user", "company_admin", "superadmin"}
            if role not in allowed_roles:
                role = "user"
            execute("UPDATE users SET role = ? WHERE id = ?", (role, user_id))
            flash("Rol bijgewerkt.")
            return redirect(url_for("users_manager"))
        if action == "reset_password":
            user_id = request.form.get("user_id", "")
            new_password = request.form.get("new_password", "").strip()
            if not new_password:
                flash("Nieuw wachtwoord ontbreekt.")
                return redirect(url_for("users_manager"))
            salt = secrets.token_hex(8)
            if is_superadmin():
                execute("UPDATE users SET salt = ?, password_hash = ? WHERE id = ?", (salt, hash_password(new_password, salt), user_id))
            else:
                execute("UPDATE users SET salt = ?, password_hash = ? WHERE id = ? AND company_id = ?", (salt, hash_password(new_password, salt), user_id, company_id))
            flash("Wachtwoord gereset.")
            return redirect(url_for("users_manager"))

    if is_superadmin():
        users = fetch_all("SELECT users.*, companies.name AS company_name FROM users JOIN companies ON companies.id = users.company_id ORDER BY users.created_at DESC")
        companies = fetch_all("SELECT * FROM companies WHERE is_active = 1 ORDER BY name")
    else:
        users = fetch_all("SELECT users.*, companies.name AS company_name FROM users JOIN companies ON companies.id = users.company_id WHERE users.company_id = ? ORDER BY users.created_at DESC", (company_id,))
        companies = []

    content = render_template_string(
        """
        <div class="grid two">
          <div class="card">
            <h2>Gebruiker toevoegen</h2>
            <form method="post">
              <input type="hidden" name="action" value="create_user">
              <input name="full_name" placeholder="Volledige naam">
              <input name="email" type="email" placeholder="E-mailadres">
              <input name="password" type="text" placeholder="Tijdelijk wachtwoord">
              {% if is_superadmin %}<select name="company_id">{% for company in companies %}<option value="{{ company['id'] }}">{{ company['name'] }}</option>{% endfor %}</select>{% endif %}
              {% if is_superadmin %}
              <select name="role">
                <option value="user">Gebruiker</option>
                <option value="company_admin">Company admin</option>
                <option value="superadmin">Superadmin</option>
              </select>
              {% else %}
              <input type="hidden" name="role" value="user">
              <div class="muted">Nieuwe accounts worden door admins altijd als gewone gebruiker aangemaakt. Alleen superadmins mogen rollen toekennen.</div>
              {% endif %}
              <button type="submit">Account maken</button>
            </form>
          </div>
          <div class="card">
            <h2>Gebruikersbeheer</h2>
            <table class="table">
              <thead><tr><th>Naam</th><th>E-mail</th><th>Bedrijf</th><th>Rol</th><th>Status</th><th>Acties</th></tr></thead>
              <tbody>
                {% for user in users %}
                <tr>
                  <td>{{ user['full_name'] }}</td>
                  <td>{{ user['email'] }}</td>
                  <td>{{ user['company_name'] }}</td>
                  <td>
                    {% if is_superadmin %}
                    <form method="post" class="inline">
                      <input type="hidden" name="action" value="change_role">
                      <input type="hidden" name="user_id" value="{{ user['id'] }}">
                      <select name="role" style="width:auto; min-width:140px;">
                        <option value="user" {{ 'selected' if user['role']=='user' else '' }}>user</option>
                        <option value="company_admin" {{ 'selected' if user['role']=='company_admin' else '' }}>company_admin</option>
                        <option value="superadmin" {{ 'selected' if user['role']=='superadmin' else '' }}>superadmin</option>
                      </select>
                      <button class="secondary" style="width:auto;">Opslaan</button>
                    </form>
                    {% else %}
                    <span class="badge">{{ user['role'] }}</span>
                    {% endif %}
                  </td>
                  <td><span class="pill {{ '' if user['is_active'] == 1 else 'off' }}">{{ 'Actief' if user['is_active'] == 1 else 'Geblokkeerd' }}</span></td>
                  <td><div class="inline"><form method="post"><input type="hidden" name="action" value="toggle_active"><input type="hidden" name="user_id" value="{{ user['id'] }}"><button class="secondary" style="width:auto;">{{ 'Blokkeer' if user['is_active'] == 1 else 'Activeer' }}</button></form><form method="post"><input type="hidden" name="action" value="reset_password"><input type="hidden" name="user_id" value="{{ user['id'] }}"><input name="new_password" type="text" placeholder="Nieuw wachtwoord" style="width:150px;"><button class="secondary" style="width:auto;">Reset</button></form>{% if user['id'] != session.get('user_id') %}<form method="post" onsubmit="return confirm('Gebruiker verwijderen?');"><input type="hidden" name="action" value="delete_user"><input type="hidden" name="user_id" value="{{ user['id'] }}"><button class="danger" style="width:auto;">Delete</button></form>{% endif %}</div></td>
                </tr>
                {% else %}<tr><td colspan="6" class="muted">Nog geen gebruikers gevonden.</td></tr>{% endfor %}
              </tbody>
            </table>
          </div>
        </div>
        """,
        users=users,
        companies=companies,
        is_superadmin=is_superadmin(),
    )
    return render_shell("Gebruikers", content)

@app.route("/feeds", methods=["GET", "POST"])
@content_admin_required
def feeds_manager() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "create_feed")
        if action == "create_feed":
            name = request.form.get("name", "").strip()
            url = request.form.get("url", "").strip()
            max_items = parse_int(request.form.get("max_items", "5"), 5, minimum=1, maximum=100)
            refresh_seconds = parse_int(request.form.get("refresh_seconds", "300"), 300, minimum=30, maximum=86400)
            is_ticker = 1 if request.form.get("is_ticker") == "on" else 0
            if not name or not url:
                flash("Naam en feed URL zijn verplicht.")
                return redirect(url_for("feeds_manager"))
            if is_ticker:
                execute("UPDATE feeds SET is_ticker = 0 WHERE company_id = ?", (company_id,))
            execute(
                "INSERT INTO feeds (id, company_id, name, url, max_items, refresh_seconds, is_active, is_ticker, created_at) VALUES (?, ?, ?, ?, ?, ?, 1, ?, ?)",
                (str(uuid.uuid4()), company_id, name, url, max_items, refresh_seconds, is_ticker, now_iso()),
            )
            flash("Feed toegevoegd.")
            return redirect(url_for("feeds_manager"))
        if action == "delete_feed":
            feed_id = request.form.get("feed_id", "")
            execute("DELETE FROM feed_items WHERE feed_id = ?", (feed_id,))
            execute("DELETE FROM feeds WHERE id = ? AND company_id = ?", (feed_id, company_id))
            flash("Feed verwijderd.")
            return redirect(url_for("feeds_manager"))
        if action == "refresh_feed":
            feed_id = request.form.get("feed_id", "")
            feed = fetch_one("SELECT * FROM feeds WHERE id = ? AND company_id = ? LIMIT 1", (feed_id, company_id))
            if feed:
                count = refresh_feed(feed)
                flash(f"Feed vernieuwd. {count} berichten opgehaald.")
            return redirect(url_for("feeds_manager"))
        if action == "toggle_feed":
            feed_id = request.form.get("feed_id", "")
            feed = fetch_one("SELECT * FROM feeds WHERE id = ? AND company_id = ? LIMIT 1", (feed_id, company_id))
            if feed:
                execute("UPDATE feeds SET is_active = ? WHERE id = ?", (0 if feed["is_active"] == 1 else 1, feed_id))
                flash("Feedstatus aangepast.")
            return redirect(url_for("feeds_manager"))
        if action == "set_ticker":
            feed_id = request.form.get("feed_id", "")
            execute("UPDATE feeds SET is_ticker = 0 WHERE company_id = ?", (company_id,))
            execute("UPDATE feeds SET is_ticker = 1 WHERE id = ? AND company_id = ?", (feed_id, company_id))
            flash("Tickerfeed ingesteld.")
            return redirect(url_for("feeds_manager"))

    feeds = fetch_all("SELECT * FROM feeds WHERE company_id = ? ORDER BY created_at DESC", (company_id,))
    feed_counts = {}
    for feed in feeds:
        rows = fetch_all("SELECT COUNT(*) AS c FROM feed_items WHERE feed_id = ?", (feed["id"],))
        feed_counts[feed["id"]] = rows[0]["c"] if rows else 0

    content = render_template_string(
        """
        <div class="grid two">
          <div class="card">
            <h2>Nieuwe feed</h2>
            <form method="post">
              <input type="hidden" name="action" value="create_feed">
              <input name="name" placeholder="Bijv. Algemeen nieuws">
              <input name="url" placeholder="Feed URL, bijv. https://example.com/rss.xml">
              <div class="inline">
                <input name="max_items" type="number" min="1" value="5" placeholder="Aantal items">
                <input name="refresh_seconds" type="number" min="30" value="300" placeholder="Refresh in sec">
              </div>
              <label class="inline"><input style="width:auto;" type="checkbox" name="is_ticker"> Gebruik als ticker</label>
              <button type="submit">Feed toevoegen</button>
            </form>
            <p class="muted">Voeg meerdere feeds toe. Kies daarna een feed als ticker voor dit bedrijf.</p>
          </div>
          <div class="card">
            <h2>Feeds</h2>
            <table class="table">
              <thead><tr><th>Naam</th><th>Items</th><th>Status</th><th>Ticker</th><th>Acties</th></tr></thead>
              <tbody>
                {% for feed in feeds %}
                <tr>
                  <td><strong>{{ feed['name'] }}</strong><br><span class="muted">{{ feed['url'] }}</span></td>
                  <td>{{ feed_counts[feed['id']] }}</td>
                  <td><span class="pill {{ '' if feed['is_active'] == 1 else 'off' }}">{{ 'Actief' if feed['is_active'] == 1 else 'Uit' }}</span></td>
                  <td>{{ 'Ja' if feed['is_ticker'] == 1 else 'Nee' }}</td>
                  <td>
                    <div class="inline">
                      <form method="post"><input type="hidden" name="action" value="refresh_feed"><input type="hidden" name="feed_id" value="{{ feed['id'] }}"><button class="secondary" style="width:auto;">Refresh</button></form>
                      <form method="post"><input type="hidden" name="action" value="toggle_feed"><input type="hidden" name="feed_id" value="{{ feed['id'] }}"><button class="secondary" style="width:auto;">Aan/uit</button></form>
                      <form method="post"><input type="hidden" name="action" value="set_ticker"><input type="hidden" name="feed_id" value="{{ feed['id'] }}"><button class="secondary" style="width:auto;">Maak ticker</button></form>
                      <form method="post" onsubmit="return confirm('Feed verwijderen?');"><input type="hidden" name="action" value="delete_feed"><input type="hidden" name="feed_id" value="{{ feed['id'] }}"><button class="danger" style="width:auto;">Delete</button></form>
                    </div>
                  </td>
                </tr>
                {% else %}
                <tr><td colspan="5" class="muted">Nog geen feeds toegevoegd.</td></tr>
                {% endfor %}
              </tbody>
            </table>
          </div>
        </div>
        """,
        feeds=feeds,
        feed_counts=feed_counts,
    )
    return render_shell("Feeds", content)

@app.route("/superadmin/companies")
@superadmin_required
def superadmin_companies() -> str:
    companies = fetch_all("SELECT * FROM companies ORDER BY created_at DESC")
    stats = []
    now_ts = time.time()
    for company in companies:
        screens = fetch_all("SELECT * FROM screens WHERE company_id = ?", (company["id"],))
        total = len(screens)
        online = sum(1 for s in screens if is_recent_heartbeat(s["last_seen"]))
        offline = total - online
        stats.append({"company": company, "total": total, "online": online, "offline": offline})
    content = render_template_string(
        """
        <div class="card">
          <h1>Superadmin bedrijven</h1>
          <p class="muted">Kies een bedrijf om de context over te nemen. Daarna werk je als superadmin binnen dat bedrijf, inclusief player installer.</p>
          <table class="table">
            <thead><tr><th>Bedrijf</th><th>Status</th><th>Schermen</th><th>Online</th><th>Offline</th><th>Acties</th></tr></thead>
            <tbody>
            {% for row in stats %}
              <tr>
                <td><strong>{{ row.company['name'] }}</strong><br><span class="muted">{{ row.company['slug'] }}</span></td>
                <td><span class="pill {{ '' if row.company['is_active'] == 1 else 'off' }}">{{ 'Actief' if row.company['is_active'] == 1 else 'Geblokkeerd' }}</span></td>
                <td>{{ row.total }}</td>
                <td>{{ row.online }}</td>
                <td>{{ row.offline }}</td>
                <td>
                  <div class="inline">
                    <form method="post" action="{{ url_for('superadmin_use_company', company_id=row.company['id']) }}">
                      <button class="secondary" style="width:auto;">Open bedrijf</button>
                    </form>
                    <form method="post" action="{{ url_for('superadmin_use_company', company_id=row.company['id']) }}">
                      <input type="hidden" name="next_page" value="player_installer">
                      <button style="width:auto;">Player installer</button>
                    </form>
                  </div>
                </td>
              </tr>
            {% endfor %}
            </tbody>
          </table>
          {% if session.get('company_view_id') %}
          <div style="margin-top:16px;">
            <form method="post" action="{{ url_for('superadmin_clear_company') }}">
              <button class="danger" style="width:auto;">Sluit bedrijfscontext</button>
            </form>
          </div>
          {% endif %}
        </div>
        """,
        stats=stats,
    )
    return render_shell("Superadmin bedrijven", content)

@app.route("/superadmin/use-company/<company_id>", methods=["POST"])
@superadmin_required
def superadmin_use_company(company_id: str):
    company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (company_id,))
    if not company:
        flash("Bedrijf niet gevonden.")
        return redirect(url_for("superadmin_companies"))
    session["company_view_id"] = company["id"]
    session["company_view_name"] = company["name"]
    flash(f"Bedrijfscontext actief: {company['name']}")
    next_page = request.form.get("next_page", "dashboard")
    if next_page not in {"dashboard", "player_installer"}:
        next_page = "dashboard"
    return redirect(url_for(next_page))

@app.route("/superadmin/clear-company", methods=["POST"])
@superadmin_required
def superadmin_clear_company():
    session.pop("company_view_id", None)
    session.pop("company_view_name", None)
    flash("Bedrijfscontext gesloten.")
    return redirect(url_for("superadmin_companies"))

@app.route("/screen-stats")
@superadmin_required
def screen_stats() -> Response | str:
    companies = fetch_all("SELECT * FROM companies ORDER BY name")
    now_ts = time.time()
    rows = []
    for company in companies:
        screens = fetch_all("SELECT * FROM screens WHERE company_id = ?", (company["id"],))
        total = len(screens)
        online = sum(1 for s in screens if is_recent_heartbeat(s["last_seen"]))
        offline = total - online
        rows.append({
            "company_name": company["name"],
            "slug": company["slug"],
            "status": "Actief" if company["is_active"] == 1 else "Geblokkeerd",
            "total": total,
            "online": online,
            "offline": offline,
        })
    if request.args.get("export") == "xlsx":
        try:
            from openpyxl import Workbook
        except Exception:
            flash("Installeer openpyxl voor Excel export: pip install openpyxl")
            return redirect(url_for("screen_stats"))
        wb = Workbook()
        ws = wb.active
        ws.title = "Schermstatistieken"
        ws.append(["Bedrijf", "Slug", "Status", "Totaal schermen", "Online", "Offline"])
        for row in rows:
            ws.append([row["company_name"], row["slug"], row["status"], row["total"], row["online"], row["offline"]])
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name="salubcast_schermstatistieken.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    content = render_template_string(
        """
        <div class="card">
          <div class="inline" style="justify-content:space-between;">
            <div>
              <h1>Schermstatistieken</h1>
              <p class="muted">Overzicht van schermen per bedrijf. Online = heartbeat in de afgelopen 120 seconden.</p>
            </div>
            <a href="{{ url_for('screen_stats', export='xlsx') }}"><button style="width:auto;">Exporteer Excel</button></a>
          </div>
          <table class="table">
            <thead><tr><th>Bedrijf</th><th>Status</th><th>Totaal</th><th>Online</th><th>Offline</th></tr></thead>
            <tbody>
            {% for row in rows %}
              <tr>
                <td><strong>{{ row.company_name }}</strong><br><span class="muted">{{ row.slug }}</span></td>
                <td><span class="pill {{ '' if row.status == 'Actief' else 'off' }}">{{ row.status }}</span></td>
                <td>{{ row.total }}</td>
                <td>{{ row.online }}</td>
                <td>{{ row.offline }}</td>
              </tr>
            {% endfor %}
            </tbody>
          </table>
        </div>
        """,
        rows=rows,
    )
    return render_shell("Schermstatistieken", content)

@app.route("/monitoring")
@superadmin_required
def monitoring_page() -> str:
    selected_company = request.args.get("company_id", "").strip()
    status_filter = request.args.get("status", "all").strip().lower()
    search = request.args.get("q", "").strip().lower()

    companies = fetch_all("SELECT * FROM companies ORDER BY name")
    company_map = {c["id"]: c for c in companies}

    query = "SELECT * FROM screens"
    params = ()
    if selected_company:
        query += " WHERE company_id = ?"
        params = (selected_company,)
    query += " ORDER BY created_at DESC"
    screens = fetch_all(query, params)

    now_ts = time.time()
    rows = []
    for screen in screens:
        company = company_map.get(screen["company_id"])
        company_name = company["name"] if company else "Onbekend bedrijf"
        company_active = bool(company and company["is_active"] == 1)

        online = False
        last_seen_label = "Nooit"
        if screen["last_seen"]:
            try:
                timestamp = iso_to_timestamp(screen["last_seen"])
                if timestamp is None:
                    raise ValueError("invalid timestamp")
                age = int(now_ts - timestamp)
                online = age < max(120, PLAYER_HEARTBEAT_SECONDS * 4)
                last_seen_label = f"{age} sec geleden"
            except Exception:
                last_seen_label = screen["last_seen"]

        ticker_feed = fetch_one(
            "SELECT name FROM feeds WHERE company_id = ? AND is_active = 1 AND is_ticker = 1 ORDER BY created_at DESC LIMIT 1",
            (screen["company_id"],),
        )

        rows.append({
            "screen_name": screen["name"],
            "location": screen["location"] or "-",
            "company_name": company_name,
            "company_active": company_active,
            "online": online,
            "status": "Online" if online else "Offline",
            "last_seen": last_seen_label,
            "last_ip": screen["device_last_ip"] or "-",
            "activated_at": screen["activated_at"] or "-",
            "ticker": ticker_feed["name"] if ticker_feed else "-",
        })

    if status_filter == "online":
        rows = [r for r in rows if r["online"]]
    elif status_filter == "offline":
        rows = [r for r in rows if not r["online"]]
    elif status_filter == "blocked":
        rows = [r for r in rows if not r["company_active"]]

    if search:
        rows = [
            r for r in rows
            if search in r["screen_name"].lower()
            or search in r["company_name"].lower()
            or search in r["location"].lower()
            or search in r["last_ip"].lower()
        ]

    content = render_template_string(
        """
        <div class="card">
          <div class="inline" style="justify-content:space-between; align-items:flex-start;">
            <div>
              <h1>Monitoring</h1>
              <p class="muted">Live overzicht van schermen, bedrijven, hartslagen en tickerstatus.</p>
            </div>
            <a href="{{ url_for('screen_stats') }}"><button style="width:auto;">Naar schermstatistieken</button></a>
          </div>

          <form method="get" class="inline" style="margin:16px 0; align-items:end;">
            <div style="min-width:220px;">
              <label class="muted">Bedrijf</label>
              <select name="company_id">
                <option value="">Alle bedrijven</option>
                {% for company in companies %}
                <option value="{{ company['id'] }}" {{ 'selected' if company['id'] == selected_company else '' }}>{{ company['name'] }}</option>
                {% endfor %}
              </select>
            </div>
            <div style="min-width:160px;">
              <label class="muted">Status</label>
              <select name="status">
                <option value="all" {{ 'selected' if status_filter == 'all' else '' }}>Alles</option>
                <option value="online" {{ 'selected' if status_filter == 'online' else '' }}>Online</option>
                <option value="offline" {{ 'selected' if status_filter == 'offline' else '' }}>Offline</option>
                <option value="blocked" {{ 'selected' if status_filter == 'blocked' else '' }}>Geblokkeerd bedrijf</option>
              </select>
            </div>
            <div style="min-width:220px;">
              <label class="muted">Zoeken</label>
              <input name="q" value="{{ search }}" placeholder="Scherm, bedrijf, locatie of IP">
            </div>
            <button type="submit" style="width:auto;">Filter</button>
            <a href="{{ url_for('monitoring_page') }}"><button type="button" class="secondary" style="width:auto;">Reset</button></a>
          </form>

          <table class="table">
            <thead>
              <tr>
                <th>Scherm</th>
                <th>Bedrijf</th>
                <th>Status</th>
                <th>Laatste heartbeat</th>
                <th>Laatste IP</th>
                <th>Geactiveerd op</th>
                <th>Ticker</th>
              </tr>
            </thead>
            <tbody>
              {% for row in rows %}
              <tr>
                <td><strong>{{ row.screen_name }}</strong><br><span class="muted">{{ row.location }}</span></td>
                <td>{{ row.company_name }}<br><span class="pill {{ '' if row.company_active else 'off' }}">{{ 'Actief' if row.company_active else 'Geblokkeerd' }}</span></td>
                <td><span class="pill {{ '' if row.online else 'off' }}">{{ row.status }}</span></td>
                <td>{{ row.last_seen }}</td>
                <td>{{ row.last_ip }}</td>
                <td>{{ row.activated_at }}</td>
                <td>{{ row.ticker }}</td>
              </tr>
              {% else %}
              <tr><td colspan="7" class="muted">Geen schermen gevonden voor deze filter.</td></tr>
              {% endfor %}
            </tbody>
          </table>
        </div>
        """,
        rows=rows,
        companies=companies,
        selected_company=selected_company,
        status_filter=status_filter,
        search=search,
    )
    return render_shell("Monitoring", content)

@app.route("/layout-studio", methods=["GET", "POST"])
@content_admin_required
def layout_studio() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        screen_id = request.form.get("screen_id", "").strip()
        orientation = request.form.get("orientation", "landscape").strip().lower()
        badge_visible = 1 if request.form.get("badge_visible") == "on" else 0
        badge_position = request.form.get("badge_position", "top-right").strip().lower()
        image_fit = request.form.get("image_fit", "contain").strip().lower()
        portrait_image_fit = request.form.get("portrait_image_fit", "contain").strip().lower()
        feed_layout = request.form.get("feed_layout", "cards").strip().lower()
        weather_city = request.form.get("weather_city", "").strip()
        insert_feed_pages = 1 if request.form.get("insert_feed_pages") == "on" else 0
        feed_page_every = max(1, int(request.form.get("feed_page_every", "3") or 3))
        feed_page_duration = max(5, int(request.form.get("feed_page_duration", "12") or 12))

        if orientation not in {"landscape", "portrait"}:
            orientation = "landscape"
        if badge_position not in {"top-right", "top-left", "bottom-right", "bottom-left"}:
            badge_position = "top-right"
        if image_fit not in {"contain", "cover"}:
            image_fit = "contain"
        if portrait_image_fit not in {"contain", "cover"}:
            portrait_image_fit = "contain"
        if feed_layout not in {"cards", "headline-list", "compact"}:
            feed_layout = "cards"

        execute(
            "UPDATE screens SET orientation = ?, badge_visible = ?, badge_position = ?, image_fit = ?, portrait_image_fit = ?, feed_layout = ?, weather_city = ?, insert_feed_pages = ?, feed_page_every = ?, feed_page_duration = ? WHERE id = ? AND company_id = ?",
            (orientation, badge_visible, badge_position, image_fit, portrait_image_fit, feed_layout, weather_city, insert_feed_pages, feed_page_every, feed_page_duration, screen_id, company_id),
        )
        flash("Layout opgeslagen.")
        return redirect(url_for("layout_studio"))

    screens = fetch_all("SELECT * FROM screens WHERE company_id = ? ORDER BY name", (company_id,))
    content = render_template_string(
        """
        <div class="card">
          <h1>Layout Studio</h1>
          <p class="muted">Stel hier per scherm portrait, badgepositie, beeldvulling en nieuws-layout in. Dit is de nette draaiknop in plaats van CSS-puzzels.</p>
        </div>
        <div class="grid">
          {% for s in screens %}
          <div class="card">
            <form method="post">
              <input type="hidden" name="screen_id" value="{{ s['id'] }}">
              <div class="inline" style="justify-content:space-between; align-items:flex-start;">
                <div>
                  <h2>{{ s['name'] }}</h2>
                  <div class="muted">{{ s['location'] or '-' }}</div>
                </div>
                <span class="pill {{ '' if s['last_seen'] else 'off' }}">{{ 'Bekend scherm' if s['last_seen'] else 'Nog niet gezien' }}</span>
              </div>
              <div class="grid cols-4" style="margin-top:14px;">
                <div><label class="muted">Orientatie</label><select name="orientation"><option value="landscape" {{ 'selected' if (s['orientation'] or 'landscape') == 'landscape' else '' }}>Landscape</option><option value="portrait" {{ 'selected' if (s['orientation'] or 'landscape') == 'portrait' else '' }}>Portrait</option></select></div>
                <div><label class="muted">Badge positie</label><select name="badge_position"><option value="top-right" {{ 'selected' if (s['badge_position'] or 'top-right') == 'top-right' else '' }}>Rechtsboven</option><option value="top-left" {{ 'selected' if (s['badge_position'] or 'top-right') == 'top-left' else '' }}>Linksboven</option><option value="bottom-right" {{ 'selected' if (s['badge_position'] or 'top-right') == 'bottom-right' else '' }}>Rechtsonder</option><option value="bottom-left" {{ 'selected' if (s['badge_position'] or 'top-right') == 'bottom-left' else '' }}>Linksonder</option></select></div>
                <div><label class="muted">Image fit landschap</label><select name="image_fit"><option value="contain" {{ 'selected' if (s['image_fit'] or 'contain') == 'contain' else '' }}>Contain</option><option value="cover" {{ 'selected' if (s['image_fit'] or 'contain') == 'cover' else '' }}>Cover</option></select></div>
                <div><label class="muted">Image fit portrait</label><select name="portrait_image_fit"><option value="contain" {{ 'selected' if (s['portrait_image_fit'] or 'contain') == 'contain' else '' }}>Contain</option><option value="cover" {{ 'selected' if (s['portrait_image_fit'] or 'contain') == 'cover' else '' }}>Cover</option></select></div>
              </div>
              <div class="grid cols-4" style="margin-top:14px;">
                <div><label class="muted">Nieuws-layout</label><select name="feed_layout"><option value="cards" {{ 'selected' if (s['feed_layout'] or 'cards') == 'cards' else '' }}>Cards</option><option value="headline-list" {{ 'selected' if (s['feed_layout'] or 'cards') == 'headline-list' else '' }}>Headline lijst</option><option value="compact" {{ 'selected' if (s['feed_layout'] or 'cards') == 'compact' else '' }}>Compact</option></select></div>
                <div><label class="muted">Weerstad</label><input name="weather_city" value="{{ s['weather_city'] or '' }}" placeholder="Bijv. Nieuwegein"><div class="muted" style="font-size:12px; margin-top:6px;">Gebruik hier altijd een stadsnaam. Schermnaam of locatie wordt niet meer gebruikt voor het weer.</div></div>
                <div><label class="muted">Nieuws elke X items</label><input type="number" min="1" name="feed_page_every" value="{{ s['feed_page_every'] or 3 }}"></div>
                <div><label class="muted">Nieuws duur (sec)</label><input type="number" min="5" name="feed_page_duration" value="{{ s['feed_page_duration'] or 12 }}"></div>
                <div style="display:flex; align-items:end; gap:12px;">
                  <label class="inline"><input style="width:auto;" type="checkbox" name="badge_visible" {{ 'checked' if s['badge_visible'] == 1 else '' }}> Badge tonen</label>
                  <label class="inline"><input style="width:auto;" type="checkbox" name="insert_feed_pages" {{ 'checked' if s['insert_feed_pages'] == 1 else '' }}> Nieuws tussendoor</label>
                </div>
              </div>
              <div class="inline" style="margin-top:14px;">
                <button type="submit" style="width:auto;">Opslaan</button>
                <a href="{{ url_for('player_page', screen=s['name'], screen_id=s['id'], token=s['token']) }}" target="_blank"><button type="button" class="secondary" style="width:auto;">Preview</button></a>
              </div>
            </form>
          </div>
          {% else %}
          <div class="card"><div class="muted">Nog geen schermen beschikbaar.</div></div>
          {% endfor %}
        </div>
        """
        , screens=screens
    )
    return render_shell("Layout Studio", content)

@app.route("/")
@login_required
def dashboard() -> str:
    company_id = current_company_id()
    company = fetch_one("SELECT * FROM companies WHERE id = ? LIMIT 1", (company_id,))
    media_count = fetch_one("SELECT COUNT(*) AS c FROM media WHERE company_id = ?", (company_id,))['c']
    screens = fetch_all("SELECT * FROM screens WHERE company_id = ? ORDER BY name", (company_id,))
    playlists = fetch_one("SELECT COUNT(*) AS c FROM playlists WHERE company_id = ?", (company_id,))['c']
    schedules = fetch_one("SELECT COUNT(*) AS c FROM schedules WHERE company_id = ?", (company_id,))['c']
    users_count = fetch_one("SELECT COUNT(*) AS c FROM users WHERE company_id = ?", (company_id,))['c']
    online = sum(1 for s in screens if is_recent_heartbeat(s['last_seen']))

    content = render_template_string(
        """
        <div class="hero">
          <div class="card">
            <div class="kicker">SaaS-ready cockpit</div>
            <h1>Welkom in SalubCast V4</h1>
            <p class="muted">Per bedrijf eigen users, schermen, playlists, branding en installerflow. Dat is nu geen proof of concept meer, maar een serieuze machine.</p>
            <div class="inline">
              {% if role == "superadmin" %}<a href="{{ url_for('player_installer') }}"><button style="width:auto;">Player installer</button></a>{% endif %}
              <a href="{{ url_for('branding_settings') }}"><button class="secondary" style="width:auto;">Branding</button></a>
            </div>
          </div>
          <div class="card">
            <h3>Huidige tenant</h3>
            <p class="muted">Bedrijf: <strong>{{ company_name }}</strong><br>Rol: <strong>{{ role }}</strong><br>Plan: <strong>{{ company['plan_name'] if company else 'starter' }}</strong><br>Billing: <strong>{{ company['billing_status'] if company else 'trial' }}</strong><br>Tagline: {{ tagline }}</p>
          </div>
        </div>
        <div class="grid cols-4">
          <div class="card"><div class="muted">Media items</div><div class="stat">{{ media_count }}</div></div>
          <div class="card"><div class="muted">Playlists</div><div class="stat">{{ playlists }}</div></div>
          <div class="card"><div class="muted">Schedules</div><div class="stat">{{ schedules }}</div></div>
          <div class="card"><div class="muted">Screens online</div><div class="stat">{{ online }}/{{ screens|length }}</div></div>
        </div>
        <div class="grid cols-4" style="margin-top:16px;"><div class="card"><div class="muted">Gebruikers in bedrijf</div><div class="stat">{{ users_count }}</div></div></div>
        <div class="grid cols-4" style="margin-top:16px;">
          {% for s in screens %}
          <div class="card">
            <h3>{{ s['name'] }}</h3>
            <div class="muted">{{ s['location'] or 'Geen locatie' }}</div>
            <div style="margin-top:12px;">{% if s['last_seen'] %}{% set online_flag = (now_ts - ts(s['last_seen'])) < 120 %}<span class="pill {{ '' if online_flag else 'off' }}">{{ 'Online' if online_flag else 'Offline' }}</span>{% else %}<span class="pill off">Nog nooit gezien</span>{% endif %}</div>
            <div class="muted" style="margin-top:10px;">Player URL: <code>{{ url_for('player_page', screen=s['name'], screen_id=s['id'], token=s['token']) }}</code></div>
          </div>
          {% endfor %}
        </div>
        """,
        company_name=current_company_name(),
        company=company,
        role=session.get('user_role'),
        media_count=media_count,
        playlists=playlists,
        schedules=schedules,
        users_count=users_count,
        online=online,
        screens=screens,
        now_ts=time.time(),
        ts=lambda iso: iso_to_timestamp(iso) or 0,
        tagline=BRAND['tagline'],
    )
    return render_shell("Dashboard", content)

@app.route("/media", methods=["GET", "POST"])
@content_admin_required
def media_library() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "upload")
        if action == "delete_media":
            media_id = request.form.get("media_id", "")
            row = fetch_one("SELECT filename FROM media WHERE id = ? AND company_id = ?", (media_id, company_id))
            if row:
                execute("DELETE FROM playlist_items WHERE media_id = ?", (media_id,))
                execute("DELETE FROM media WHERE id = ? AND company_id = ?", (media_id, company_id))
                path = UPLOAD_DIR / row['filename']
                if path.exists():
                    path.unlink()
                flash("Media verwijderd.")
            return redirect(url_for("media_library"))
        file = request.files.get("file")
        title = request.form.get("title", "").strip() or "Untitled"
        duration = parse_int(request.form.get("duration_seconds", "10"), 10, minimum=1, maximum=86400)
        if not file or not file.filename:
            flash("Kies eerst een bestand.")
            return redirect(url_for("media_library"))
        valid, message = validate_upload_file(file)
        if not valid:
            flash(message)
            return redirect(url_for("media_library"))
        safe_name = secure_filename(file.filename)
        unique_name = f"{uuid.uuid4()}_{safe_name}"
        file.save(UPLOAD_DIR / unique_name)
        media_id = str(uuid.uuid4())
        execute("INSERT INTO media (id, company_id, title, filename, mimetype, duration_seconds, uploaded_at) VALUES (?, ?, ?, ?, ?, ?, ?)", (media_id, company_id, title, unique_name, get_mimetype(safe_name), duration, now_iso()))
        log_event(actor_label(), "media_uploaded", "media", media_id, title, company_id)
        flash("Media toegevoegd.")
        return redirect(url_for("media_library"))
    items = fetch_all("SELECT * FROM media WHERE company_id = ? ORDER BY uploaded_at DESC", (company_id,))
    safe_items = []
    for item in items:
        row = dict(item)
        filename = str(row.get("filename") or "")
        mimetype = str(row.get("mimetype") or "")
        if not mimetype and "." in filename:
            mimetype = get_mimetype(filename)
        row["filename"] = filename
        row["mimetype"] = mimetype or "application/octet-stream"
        row["title"] = str(row.get("title") or "Untitled")
        row["duration_seconds"] = row.get("duration_seconds") or 10
        safe_items.append(row)
    items = safe_items
    content = render_template_string(
        """
        <div class="grid two"><div class="card"><h2>Upload nieuwe media</h2><form method="post" enctype="multipart/form-data"><input type="hidden" name="action" value="upload"><input name="title" placeholder="Titel"><input name="duration_seconds" type="number" min="1" value="10" placeholder="Duur in seconden"><input name="file" type="file" accept="image/*,video/*,.pdf"><button type="submit">Uploaden</button></form></div><div class="card"><h2>Bibliotheek</h2><table class="table"><thead><tr><th>Preview</th><th>Titel</th><th>Type</th><th>Duur</th><th>Actie</th></tr></thead><tbody>{% for item in items %}<tr><td>{% if item['mimetype'].startswith('image/') %}<img class="media-preview" src="{{ url_for('uploaded_file', filename=item['filename']) }}">{% elif item['mimetype'].startswith('video/') %}[VID]{% elif item['mimetype'] == 'application/pdf' %}[PDF]{% else %}-{% endif %}</td><td>{{ item['title'] }}</td><td>{{ item['mimetype'] }}</td><td>{{ item['duration_seconds'] }}s</td><td><form method="post" onsubmit="return confirm('Media verwijderen?');"><input type="hidden" name="action" value="delete_media"><input type="hidden" name="media_id" value="{{ item['id'] }}"><button type="submit" class="danger">Delete</button></form></td></tr>{% else %}<tr><td colspan="5" class="muted">Nog geen media aanwezig.</td></tr>{% endfor %}</tbody></table></div></div>
        """,
        items=items,
    )
    return render_shell("Media", content)

@app.route("/playlists", methods=["GET", "POST"])
@content_admin_required
def playlist_manager() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "create_playlist":
            name = request.form.get("name", "").strip()
            if name:
                execute("INSERT INTO playlists (id, company_id, name, created_at) VALUES (?, ?, ?, ?)", (str(uuid.uuid4()), company_id, name, now_iso()))
                flash("Playlist aangemaakt.")
        elif action == "delete_playlist":
            playlist_id = request.form.get("playlist_id", "")
            playlist = fetch_company_row("playlists", playlist_id, company_id)
            if not playlist:
                flash("Playlist hoort niet bij dit bedrijf.")
                return redirect(url_for("playlist_manager"))
            execute("DELETE FROM playlist_items WHERE playlist_id = ?", (playlist_id,))
            execute("DELETE FROM schedules WHERE playlist_id = ? AND company_id = ?", (playlist_id, company_id))
            execute("DELETE FROM playlists WHERE id = ? AND company_id = ?", (playlist_id, company_id))
            flash("Playlist verwijderd.")
        elif action == "add_item":
            playlist_id = request.form.get("playlist_id", "")
            media_id = request.form.get("media_id", "")
            playlist = fetch_company_row("playlists", playlist_id, company_id)
            media = fetch_company_row("media", media_id, company_id)
            if not playlist or not media:
                flash("Playlist of media hoort niet bij dit bedrijf.")
                return redirect(url_for("playlist_manager"))
            current = fetch_one("SELECT COALESCE(MAX(sort_order), 0) AS max_sort FROM playlist_items WHERE playlist_id = ?", (playlist_id,))['max_sort']
            execute("INSERT INTO playlist_items (id, playlist_id, media_id, sort_order) VALUES (?, ?, ?, ?)", (str(uuid.uuid4()), playlist_id, media_id, int(current) + 1))
            flash("Media toegevoegd aan playlist.")
        elif action == "remove_item":
            item_id = request.form.get("item_id", "")
            item = fetch_one("SELECT pi.id FROM playlist_items pi JOIN playlists p ON p.id = pi.playlist_id WHERE pi.id = ? AND p.company_id = ? LIMIT 1", (item_id, company_id))
            if item:
                execute("DELETE FROM playlist_items WHERE id = ?", (item_id,))
                flash("Item verwijderd uit playlist.")
        elif action == "move_up":
            item_id = request.form.get("item_id", "")
            item = fetch_one("SELECT pi.* FROM playlist_items pi JOIN playlists p ON p.id = pi.playlist_id WHERE pi.id = ? AND p.company_id = ? LIMIT 1", (item_id, company_id))
            if item:
                prev_item = fetch_one("SELECT * FROM playlist_items WHERE playlist_id = ? AND sort_order < ? ORDER BY sort_order DESC LIMIT 1", (item['playlist_id'], item['sort_order']))
                if prev_item:
                    execute("UPDATE playlist_items SET sort_order = ? WHERE id = ?", (prev_item['sort_order'], item['id']))
                    execute("UPDATE playlist_items SET sort_order = ? WHERE id = ?", (item['sort_order'], prev_item['id']))
        elif action == "move_down":
            item_id = request.form.get("item_id", "")
            item = fetch_one("SELECT pi.* FROM playlist_items pi JOIN playlists p ON p.id = pi.playlist_id WHERE pi.id = ? AND p.company_id = ? LIMIT 1", (item_id, company_id))
            if item:
                next_item = fetch_one("SELECT * FROM playlist_items WHERE playlist_id = ? AND sort_order > ? ORDER BY sort_order ASC LIMIT 1", (item['playlist_id'], item['sort_order']))
                if next_item:
                    execute("UPDATE playlist_items SET sort_order = ? WHERE id = ?", (next_item['sort_order'], item['id']))
                    execute("UPDATE playlist_items SET sort_order = ? WHERE id = ?", (item['sort_order'], next_item['id']))
        return redirect(url_for("playlist_manager"))
    playlists = fetch_all("SELECT * FROM playlists WHERE company_id = ? ORDER BY created_at DESC", (company_id,))
    media_items = fetch_all("SELECT * FROM media WHERE company_id = ? ORDER BY uploaded_at DESC", (company_id,))
    playlist_map = {}
    for p in playlists:
        playlist_map[p['id']] = fetch_all("SELECT pi.id, pi.sort_order, m.title, m.mimetype FROM playlist_items pi JOIN media m ON m.id = pi.media_id WHERE pi.playlist_id = ? ORDER BY pi.sort_order ASC", (p['id'],))
    content = render_template_string(
        """
<div class="grid two"><div class="card"><h2>Playlist builder</h2><form method="post"><input type="hidden" name="action" value="create_playlist"><input name="name" placeholder="Bijv. Reception Morning Loop"><button type="submit">Nieuwe playlist</button></form><hr style="border-color:#334155; opacity:.35; margin:18px 0;"><form method="post"><input type="hidden" name="action" value="add_item"><select name="playlist_id">{% for p in playlists %}<option value="{{ p['id'] }}">{{ p['name'] }}</option>{% endfor %}</select><select name="media_id">{% for m in media_items %}<option value="{{ m['id'] }}">{{ m['title'] }} ({{ m['mimetype'] }})</option>{% endfor %}</select><button type="submit">Media toevoegen</button></form></div><div class="grid">{% for p in playlists %}<div class="card"><div class="inline" style="justify-content:space-between; align-items:start;"><div><h3>{{ p['name'] }}</h3><div class="muted">Sleutel de volgorde met omhoog/omlaag-knoppen.</div></div><form method="post" onsubmit="return confirm('Playlist verwijderen?');"><input type="hidden" name="action" value="delete_playlist"><input type="hidden" name="playlist_id" value="{{ p['id'] }}"><button class="danger" style="width:auto;">Delete playlist</button></form></div><div class="dropzone" style="margin-top:12px;">{% for item in playlist_map[p['id']] %}<div class="playlist-row"><div><span class="badge">#{{ item['sort_order'] }}</span></div><div><strong>{{ item['title'] }}</strong><br><span class="muted">{{ item['mimetype'] }}</span></div><div class="inline" style="justify-content:flex-end;"><form method="post"><input type="hidden" name="action" value="move_up"><input type="hidden" name="item_id" value="{{ item['id'] }}"><button class="secondary" style="width:auto;">Omhoog</button></form><form method="post"><input type="hidden" name="action" value="move_down"><input type="hidden" name="item_id" value="{{ item['id'] }}"><button class="secondary" style="width:auto;">Omlaag</button></form><form method="post"><input type="hidden" name="action" value="remove_item"><input type="hidden" name="item_id" value="{{ item['id'] }}"><button class="danger" style="width:auto;">Verwijder</button></form></div></div>{% else %}<div class="muted">Lege playlist. Tijd om pixels los te laten.</div>{% endfor %}</div></div>{% endfor %}</div></div>
        """,
        playlists=playlists,
        media_items=media_items,
        playlist_map=playlist_map,
    )
    return render_shell("Playlists", content)

@app.route("/screens", methods=["GET", "POST"])
@admin_required
def screens_manager() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "create":
            name = request.form.get("name", "").strip()
            location = request.form.get("location", "").strip()
            orientation = request.form.get("orientation", "landscape").strip().lower()
            insert_feed_pages = 1 if request.form.get("insert_feed_pages") == "on" else 0
            feed_page_every = parse_int(request.form.get("feed_page_every", "3"), 3, minimum=1, maximum=1000)
            feed_page_duration = parse_int(request.form.get("feed_page_duration", "12"), 12, minimum=5, maximum=86400)
            if orientation not in {"landscape", "portrait"}:
                orientation = "landscape"
            if name:
                new_screen_id = str(uuid.uuid4())
                execute(
                    "INSERT INTO screens (id, company_id, name, location, orientation, insert_feed_pages, feed_page_every, feed_page_duration, token, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (new_screen_id, company_id, name, location, orientation, insert_feed_pages, feed_page_every, feed_page_duration, str(uuid.uuid4()), now_iso()),
                )
                log_event(actor_label(), "screen_created", "screen", new_screen_id, name, company_id)
                flash("Scherm toegevoegd.")
        elif action == "save_settings":
            screen_id = request.form.get("screen_id", "")
            orientation = request.form.get("orientation", "landscape").strip().lower()
            insert_feed_pages = 1 if request.form.get("insert_feed_pages") == "on" else 0
            feed_page_every = parse_int(request.form.get("feed_page_every", "3"), 3, minimum=1, maximum=1000)
            feed_page_duration = parse_int(request.form.get("feed_page_duration", "12"), 12, minimum=5, maximum=86400)
            if orientation not in {"landscape", "portrait"}:
                orientation = "landscape"
            execute(
                "UPDATE screens SET orientation = ?, insert_feed_pages = ?, feed_page_every = ?, feed_page_duration = ? WHERE id = ? AND company_id = ?",
                (orientation, insert_feed_pages, feed_page_every, feed_page_duration, screen_id, company_id),
            )
            flash("Scherminstellingen bijgewerkt.")
        elif action == "delete_screen":
            screen_id = request.form.get("screen_id", "")
            execute("DELETE FROM schedules WHERE screen_id = ? AND company_id = ?", (screen_id, company_id))
            execute("DELETE FROM screens WHERE id = ? AND company_id = ?", (screen_id, company_id))
            log_event(actor_label(), "screen_deleted", "screen", screen_id, "", company_id)
            flash("Scherm verwijderd.")
        return redirect(url_for("screens_manager"))
    screens = fetch_all("SELECT * FROM screens WHERE company_id = ? ORDER BY created_at DESC", (company_id,))
    content = render_template_string(
        """
        <div class="grid two">
          <div class="card">
            <h2>Nieuw scherm</h2>
            <form method="post">
              <input type="hidden" name="action" value="create">
              <input name="name" placeholder="Bijv. Kantine scherm">
              <input name="location" placeholder="Locatie">
              <select name="orientation">
                <option value="landscape">Landscape</option>
                <option value="portrait">Portrait / verticaal</option>
              </select>
              <label class="inline"><input style="width:auto;" type="checkbox" name="insert_feed_pages"> Toon nieuws-pagina's tussen content</label>
              <div class="inline">
                <input type="number" min="1" name="feed_page_every" value="3" placeholder="Elke X items">
                <input type="number" min="5" name="feed_page_duration" value="12" placeholder="Nieuws-pagina duur">
              </div>
              <button type="submit">Scherm opslaan</button>
            </form>
          </div>
          <div class="card">
            <h2>Schermen</h2>
            <table class="table">
              <thead><tr><th>Naam</th><th>Locatie</th><th>Status</th><th>Weergave</th><th>Nieuws</th><th>Activatie</th><th>Actie</th></tr></thead>
              <tbody>
              {% for s in screens %}
                {% set online_flag = s['last_seen'] and (now_ts - ts(s['last_seen'])) < 120 %}
                <tr>
                  <td><strong>{{ s['name'] }}</strong></td>
                  <td>{{ s['location'] }}</td>
                  <td><span class="pill {{ '' if online_flag else 'off' }}">{{ 'Online' if online_flag else 'Offline' }}</span></td>
                  <td>
                    <form method="post" class="inline">
                      <input type="hidden" name="action" value="save_settings">
                      <input type="hidden" name="screen_id" value="{{ s['id'] }}">
                      <select name="orientation" style="width:auto;">
                        <option value="landscape" {{ 'selected' if (s['orientation'] or 'landscape') == 'landscape' else '' }}>Landscape</option>
                        <option value="portrait" {{ 'selected' if (s['orientation'] or 'landscape') == 'portrait' else '' }}>Portrait</option>
                      </select>
                  </td>
                  <td>
                      <label class="inline"><input style="width:auto;" type="checkbox" name="insert_feed_pages" {{ 'checked' if s['insert_feed_pages'] == 1 else '' }}> Aan</label>
                      <input type="number" min="1" name="feed_page_every" value="{{ s['feed_page_every'] or 3 }}" style="width:72px;">
                      <input type="number" min="5" name="feed_page_duration" value="{{ s['feed_page_duration'] or 12 }}" style="width:72px;">
                      <button class="secondary" style="width:auto;">Opslaan</button>
                    </form>
                  </td>
                  <td>{{ s['activation_code'] or '-' }}</td>
                  <td>
                    <form method="post" onsubmit="return confirm('Scherm verwijderen?');">
                      <input type="hidden" name="action" value="delete_screen">
                      <input type="hidden" name="screen_id" value="{{ s['id'] }}">
                      <button class="danger" style="width:auto;">Delete</button>
                    </form>
                  </td>
                </tr>
              {% endfor %}
              </tbody>
            </table>
          </div>
        </div>
        """,
        screens=screens,
        now_ts=time.time(),
        ts=lambda iso: iso_to_timestamp(iso) or 0,
    )
    return render_shell("Screens", content)

@app.route("/schedules", methods=["GET", "POST"])
@admin_required
def schedules_manager() -> str:
    company_id = current_company_id()
    if request.method == "POST":
        action = request.form.get("action", "create")
        schedule_id = request.form.get("schedule_id", "").strip()
        if action in {"create", "update"}:
            screen_id = request.form.get("screen_id", "")
            playlist_id = request.form.get("playlist_id", "")
            start_time = request.form.get("start_time", "08:00")
            end_time = request.form.get("end_time", "18:00")
            days = [int(d) for d in request.form.getlist("days")]
            priority = parse_int(request.form.get("priority", "100"), 100, minimum=1, maximum=999)
            if not days:
                flash("Kies minimaal een dag.")
                return redirect(url_for("schedules_manager"))
            if not fetch_company_row("screens", screen_id, company_id):
                flash("Gekozen scherm hoort niet bij dit bedrijf.")
                return redirect(url_for("schedules_manager"))
            if not fetch_company_row("playlists", playlist_id, company_id):
                flash("Gekozen playlist hoort niet bij dit bedrijf.")
                return redirect(url_for("schedules_manager"))
            conflicts = schedule_conflicts(company_id, screen_id, days, start_time, end_time, schedule_id if action == 'update' else None)
            if conflicts:
                flash(f"Schema conflict met {len(conflicts)} bestaand(e) schema('s). Pas tijden of dagen aan.")
                return redirect(url_for("schedules_manager"))
            if action == "create":
                new_id = str(uuid.uuid4())
                execute("INSERT INTO schedules (id, company_id, screen_id, playlist_id, start_time, end_time, days_json, priority, active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)", (new_id, company_id, screen_id, playlist_id, start_time, end_time, json.dumps(days), priority))
                log_event(actor_label(), "schedule_created", "schedule", new_id, f"{start_time}-{end_time} priority {priority}", company_id)
                flash("Schema opgeslagen.")
            else:
                execute("UPDATE schedules SET screen_id = ?, playlist_id = ?, start_time = ?, end_time = ?, days_json = ?, priority = ? WHERE id = ? AND company_id = ?", (screen_id, playlist_id, start_time, end_time, json.dumps(days), priority, schedule_id, company_id))
                log_event(actor_label(), "schedule_updated", "schedule", schedule_id, f"{start_time}-{end_time} priority {priority}", company_id)
                flash("Schema bijgewerkt.")
        elif action == "delete":
            execute("DELETE FROM schedules WHERE id = ? AND company_id = ?", (schedule_id, company_id))
            log_event(actor_label(), "schedule_deleted", "schedule", schedule_id, "", company_id)
            flash("Schema verwijderd.")
        elif action == "toggle":
            sch = fetch_one("SELECT * FROM schedules WHERE id = ? AND company_id = ? LIMIT 1", (schedule_id, company_id))
            if sch:
                execute("UPDATE schedules SET active = ? WHERE id = ? AND company_id = ?", (0 if sch['active'] == 1 else 1, schedule_id, company_id))
                log_event(actor_label(), "schedule_toggled", "schedule", schedule_id, f"active={0 if sch['active'] == 1 else 1}", company_id)
                flash("Schemastatus aangepast.")
        return redirect(url_for("schedules_manager"))
    screens = fetch_all("SELECT * FROM screens WHERE company_id = ? ORDER BY name", (company_id,))
    playlists = fetch_all("SELECT * FROM playlists WHERE company_id = ? ORDER BY name", (company_id,))
    schedules = fetch_all("SELECT sc.id, sc.start_time, sc.end_time, sc.days_json, sc.priority, sc.active, s.name AS screen_name, p.name AS playlist_name FROM schedules sc JOIN screens s ON s.id = sc.screen_id JOIN playlists p ON p.id = sc.playlist_id WHERE sc.company_id = ? ORDER BY s.name, sc.priority ASC, sc.start_time", (company_id,))
    day_names = ["Ma", "Di", "Wo", "Do", "Vr", "Za", "Zo"]
    content = render_template_string(
        """
        <div class="grid two"><div class="card"><h2>Nieuw schema</h2><form method="post"><input type="hidden" name="action" value="create"><select name="screen_id">{% for s in screens %}<option value="{{ s['id'] }}">{{ s['name'] }}</option>{% endfor %}</select><select name="playlist_id">{% for p in playlists %}<option value="{{ p['id'] }}">{{ p['name'] }}</option>{% endfor %}</select><div class="inline"><input type="time" name="start_time" value="08:00"><input type="time" name="end_time" value="18:00"></div><input type="number" name="priority" value="100" min="1" max="999" placeholder="Prioriteit"><div class="dropzone">{% for idx, dn in day_names %}<label class="inline" style="width:48%; margin:4px 0;"><input style="width:auto;" type="checkbox" name="days" value="{{ idx }}" checked> {{ dn }}</label>{% endfor %}</div><button type="submit">Schema opslaan</button></form><p class="muted">Overnacht-schema's en conflictcontrole zijn ingebouwd.</p></div><div class="card"><h2>Actieve schema's</h2><table class="table"><thead><tr><th>Scherm</th><th>Playlist</th><th>Tijd</th><th>Prio</th><th>Dagen</th><th>Status</th><th>Acties</th></tr></thead><tbody>{% for item in schedules %}<tr><td>{{ item['screen_name'] }}</td><td>{{ item['playlist_name'] }}</td><td>{{ item['start_time'] }} - {{ item['end_time'] }}</td><td>{{ item['priority'] }}</td><td>{% for d in from_json(item['days_json']) %}{{ short_day(d) }} {% endfor %}</td><td><span class="pill {{ '' if item['active'] == 1 else 'off' }}">{{ 'Actief' if item['active'] == 1 else 'Uit' }}</span></td><td><div class="inline"><form method="post"><input type="hidden" name="action" value="toggle"><input type="hidden" name="schedule_id" value="{{ item['id'] }}"><button class="secondary" style="width:auto;">Aan/uit</button></form><form method="post" onsubmit="return confirm('Schema verwijderen?');"><input type="hidden" name="action" value="delete"><input type="hidden" name="schedule_id" value="{{ item['id'] }}"><button class="danger" style="width:auto;">Delete</button></form></div></td></tr>{% else %}<tr><td colspan="7" class="muted">Nog geen schema's.</td></tr>{% endfor %}</tbody></table></div></div>
        """,
        screens=screens,
        playlists=playlists,
        schedules=schedules,
        day_names=list(enumerate(day_names)),
        from_json=json.loads,
        short_day=lambda idx: day_names[idx],
    )
    return render_shell("Schedules", content)

@app.route("/audit-logs")
@superadmin_required
def audit_logs_page() -> str:
    company_id = current_company_id()
    logs = fetch_all("SELECT * FROM audit_logs WHERE company_id IS NULL OR company_id = ? ORDER BY created_at DESC LIMIT 250", (company_id,))
    content = render_template_string("""
        <div class="card"><h1>Audit logs</h1><table class="table"><thead><tr><th>Tijd</th><th>Actor</th><th>Actie</th><th>Type</th><th>Target</th><th>Details</th></tr></thead><tbody>{% for log in logs %}<tr><td>{{ log['created_at'] }}</td><td>{{ log['actor'] }}</td><td>{{ log['action'] }}</td><td>{{ log['target_type'] }}</td><td>{{ log['target_id'] }}</td><td>{{ log['details'] }}</td></tr>{% else %}<tr><td colspan="6" class="muted">Nog geen logs.</td></tr>{% endfor %}</tbody></table></div>
    """, logs=logs)
    return render_shell("Audit logs", content)
